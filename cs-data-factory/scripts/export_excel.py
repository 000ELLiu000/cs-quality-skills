#!/usr/bin/env python3
"""
export_excel.py
---------------
Phase 2 可选导出：将 final.jsonl 转换为带格式的 Excel 文件（final.xlsx）。

功能：
- 读取指定 run-dir 下的 final.jsonl
- 输出同目录下的 final.xlsx（深蓝表头 + 交替行底色 + 首行冻结 + 自动筛选）
- 依赖：openpyxl（pip install openpyxl）

用法：
  python scripts/export_excel.py --run-dir data/runs/run_YYYYMMDD_NNN

选项：
  --run-dir    运行目录（必须已包含 final.jsonl；支持 data/runs/ 自动补全）
  --out        输出文件路径（可选，默认为 run-dir/final.xlsx）
"""

import json
import sys
import argparse
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("❌ 缺少依赖：pip install openpyxl", file=sys.stderr)
    sys.exit(1)

SKILL_ROOT = Path(__file__).parent.parent

HEADERS = ["会话id", "消息时间", "历史消息", "话术回复",
           "话术关联动作", "话术共识动作", "所有信号", "外呼信号"]
COL_WIDTHS = [28, 18, 45, 45, 20, 20, 50, 25]


def export(jsonl_path: Path, out_path: Path) -> int:
    rows = []
    with open(jsonl_path, encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if line:
                rows.append(json.loads(line))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "生产数据"

    # 表头
    header_fill = PatternFill("solid", start_color="366092", end_color="366092")
    header_font = Font(bold=True, color="FFFFFF", name="Arial", size=10)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col_idx, h in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
    ws.row_dimensions[1].height = 30

    # 数据行
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    alt_fill = PatternFill("solid", start_color="EBF1F8", end_color="EBF1F8")

    for row_idx, row in enumerate(rows, 2):
        use_alt = (row_idx % 2 == 0)
        for col_idx, h in enumerate(HEADERS, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=row.get(h, ""))
            cell.font = Font(name="Arial", size=9)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
            if use_alt:
                cell.fill = alt_fill

    # 列宽、冻结、筛选
    for i, w in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}1"

    wb.save(out_path)
    return len(rows)


def main():
    parser = argparse.ArgumentParser(description="将 final.jsonl 导出为 Excel")
    parser.add_argument("--run-dir", required=True,
                        help="运行目录（支持 data/runs/ 自动补全）")
    parser.add_argument("--out", default=None,
                        help="输出文件路径（默认：run-dir/final.xlsx）")
    args = parser.parse_args()

    raw = args.run_dir
    if not Path(raw).is_absolute() and not raw.startswith("data/runs"):
        raw = f"data/runs/{raw}"
    run_dir = SKILL_ROOT / raw

    jsonl_path = run_dir / "final.jsonl"
    if not jsonl_path.exists():
        print(f"❌ 找不到 final.jsonl：{jsonl_path}", file=sys.stderr)
        sys.exit(1)

    out_path = Path(args.out) if args.out else run_dir / "final.xlsx"
    n = export(jsonl_path, out_path)
    print(f"✅ 已导出 Excel：{out_path}（{n} 行数据）")


if __name__ == "__main__":
    main()
